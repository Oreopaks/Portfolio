"""Тесты фич владельца: алерты, /stats, подсказка перебить цену, /export."""
from mvp7_price_scout import store, handlers, alerts
from mvp7_price_scout.normalize import Product


def _seed_base(conn, our=150000):
    b = Product(shop="di-park", title="Apple iPhone 17 Pro Max 256Gb", price=our,
                source_type="base", fetched_at="2026-06-29T16:00:00")
    store.replace_shop(conn, "di-park", [b])
    store.link_base_self(conn)
    return store.base_catalog(conn)[0]["id"]


def test_alert_on_new_undercut(tmp_path):
    conn = store.connect(tmp_path / "p.db")
    bid = _seed_base(conn, our=150000)
    # прошлый прогон: конкурент был ДОРОЖЕ нас (151000)
    prev = Product(shop="sr57", title="x", price=151000)
    prev.dipark_id = bid
    store.append_price_log(conn, [prev], "2026-06-29T10:00:00")
    # текущий срез: конкурент стал ДЕШЕВЛЕ нас (144000)
    now = Product(shop="sr57", title="Apple iPhone 17 Pro Max 256Gb", price=144000)
    now.dipark_id = bid
    store.replace_shop(conn, "sr57", [now])

    msgs = alerts.compute_alerts(conn, "2026-06-29T12:00:00")
    assert any("дешевле нас" in m for m in msgs)


def test_alert_dedups_same_shop(tmp_path):
    """Две строки одного магазина к одному товару -> ОДИН алерт (не дубль «Marshall ×2»)."""
    conn = store.connect(tmp_path / "p.db")
    bid = _seed_base(conn, our=150000)
    prev = Product(shop="sr57", title="x", price=151000)   # раньше был дороже нас
    prev.dipark_id = bid
    store.append_price_log(conn, [prev], "2026-06-29T10:00:00")
    rows = [Product(shop="sr57", title="iPhone 17 Pro Max 256 A", price=144000),
            Product(shop="sr57", title="iPhone 17 Pro Max 256 B", price=145000)]
    for r in rows:
        r.dipark_id = bid
    store.replace_shop(conn, "sr57", rows)
    msgs = alerts.compute_alerts(conn, "2026-06-29T12:00:00")
    assert sum("дешевле нас" in m for m in msgs) == 1


def test_no_alert_without_history(tmp_path):
    conn = store.connect(tmp_path / "p.db")
    bid = _seed_base(conn)
    now = Product(shop="sr57", title="iPhone 17 Pro Max 256", price=100)
    now.dipark_id = bid
    store.replace_shop(conn, "sr57", [now])
    assert alerts.compute_alerts(conn, "2026-06-29T12:00:00") == []   # нет прошлого прогона


def test_stats_position(tmp_path):
    conn = store.connect(tmp_path / "p.db")
    bid = _seed_base(conn, our=150000)
    c = Product(shop="sr57", title="iPhone 17 Pro Max 256", price=144000)  # мы дороже
    c.dipark_id = bid
    store.replace_shop(conn, "sr57", [c])
    mp = store.market_position(conn)
    assert mp["total"] == 1 and mp["pricier"] == 1 and mp["cheaper"] == 0
    assert "Ты дороже всех:   1" in handlers.render_stats(conn)


def test_reprice_hint_when_we_are_pricier(tmp_path):
    conn = store.connect(tmp_path / "p.db")
    bid = _seed_base(conn, our=150000)
    c = Product(shop="sr57", title="iPhone 17 Pro Max 256", price=144000)
    c.dipark_id = bid
    store.replace_shop(conn, "sr57", [c])
    out = handlers.render_comparison(conn, store.base_by_id(conn, bid))
    assert "🎯 Поставь 143 900 ₽ → станешь дешевле всех" in out


def test_export_xlsx(tmp_path):
    conn = store.connect(tmp_path / "p.db")
    bid = _seed_base(conn, our=150000)
    c = Product(shop="sr57", title="iPhone 17 Pro Max 256", price=144000)
    c.dipark_id = bid
    store.replace_shop(conn, "sr57", [c])
    path = handlers.build_export_xlsx(conn)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["Сравнение цен"]
    headers = [c.value for c in ws[1]]
    assert "Товар" in headers and "sr57.ru" in headers and "Статус" in headers
    row2 = [c.value for c in ws[2]]
    assert 150000 in row2 and 144000 in row2          # наша и конкурентная цена
    assert "🔴 дороже всех" in row2                    # статус
    assert "Сводка" in wb.sheetnames
